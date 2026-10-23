"""显式保留的兼容本地基线解析器。"""
import io
import re
import zipfile
from pathlib import Path
from typing import BinaryIO, Protocol

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.exc import PackageNotFoundError
from pypdf import PdfReader

LOCAL_PARSER_ID = "local-baseline@2"


class ParserAdapter(Protocol):
    """MineU等未来解析引擎通过相同的来源区间契约接入。"""

    version: str

    def parse(self, name: str, data: bytes | BinaryIO, revision: str) -> dict: ...


class LocalParserAdapter:
    version = LOCAL_PARSER_ID

    def parse(self, name, data, revision):
        return parse_local(name, data, revision)


def _stream(data):
    return data if hasattr(data, "read") else io.BytesIO(data)


def parse_local(name, data, revision):
    ext, blocks = Path(name).suffix.lower(), []

    def add(text, kind="text", page=None, section=None, rows=None, slide=None, shape_id=None,
            shape_name=None, position=None):
        if text.strip():
            blocks.append({"id": f"b{len(blocks)}", "type": kind, "text": text, "page": page,
                           "section_path": section or [], "rows": rows, "slide": slide,
                           "shape_id": shape_id, "shape_name": shape_name, "position": position})

    if ext in {".txt", ".md"}:
        source = data.read() if hasattr(data, "read") else data
        raw = source.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
        section = []
        for group in re.split(r"\n\s*\n", raw):
            if group.startswith("#"):
                section = [group.splitlines()[0].lstrip("# ")]
            rows = group.splitlines()
            table = len(rows) > 1 and all("|" in r for r in rows)
            add(group, "table" if table else "text", section=section, rows=rows if table else None)
    elif ext == ".docx":
        document, section = Document(_stream(data)), []
        for child in document.element.body:
            if child.tag.endswith("}p"):
                p = Paragraph(child, document)
                if p.style and p.style.name.startswith("Heading"):
                    section = [p.text]
                add(p.text, section=section)
            elif child.tag.endswith("}tbl"):
                rows = [" | ".join(cell.text for cell in row.cells) for row in Table(child, document).rows]
                add("\n".join(rows), "table", section=section, rows=rows)
    elif ext == ".pdf":
        reader = PdfReader(_stream(data))
        if reader.is_encrypted:
            raise ValueError("暂不支持加密PDF")
        for page, content in enumerate(reader.pages, 1):
            add(content.extract_text() or "", page=page)
    elif ext == ".pptx":
        try:
            presentation = Presentation(_stream(data))
        except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise ValueError("PPTX文件无效或无法解析") from exc
        for slide_number, slide in enumerate(presentation.slides, 1):
            title_shape = slide.shapes.title
            section = [title_shape.text.strip()] if title_shape and title_shape.text.strip() else []
            for shape in slide.shapes:
                shape_id, shape_name = str(shape.shape_id), shape.name
                position = {"left": int(shape.left), "top": int(shape.top),
                            "width": int(shape.width), "height": int(shape.height)}
                if shape.has_table:
                    rows = [" | ".join(cell.text for cell in row.cells) for row in shape.table.rows]
                    add("\n".join(rows), "table", section=section, rows=rows, slide=slide_number,
                        shape_id=shape_id, shape_name=shape_name, position=position)
                elif shape.has_text_frame:
                    add(shape.text_frame.text, section=section, slide=slide_number,
                        shape_id=shape_id, shape_name=shape_name, position=position)
    else:
        raise ValueError("支持TXT/MD/DOCX/文本PDF/PPTX；旧版PPT请先另存为PPTX")
    if not blocks:
        raise ValueError("未提取到文本；扫描件请安装并配置Docling PDF模型资产及OCR")
    result = {"schema_version": 1, "revision": revision, "parser": LOCAL_PARSER_ID,
              "blocks": blocks, "warnings": ["兼容本地基线不保证DOCX页码与PDF表格结构",
                                             "PPTX图片中的文字不执行OCR"]}
    if ext == ".pptx":
        result["slide_size"] = {"width": int(presentation.slide_width), "height": int(presentation.slide_height)}
    return result
