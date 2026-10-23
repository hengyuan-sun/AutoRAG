import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from raglab.chunking import presets
from raglab.config import Settings
from raglab.parsing import parse_local
from raglab.service import Lab


def test_presets_explain_strategy_specific_length_semantics():
    by_id = {preset["id"]: preset for preset in presets()}
    assert set(by_id) == {"recursive", "window", "heading", "faq", "table"}
    assert by_id["recursive"]["length_label"] == "单块最大长度（字符）"
    assert "段落和句子边界" in by_id["recursive"]["length_semantics"]
    assert "超长" in by_id["recursive"]["length_semantics"]
    assert by_id["window"]["length_label"] == "固定窗口长度（字符）"
    assert by_id["table"]["overlap_supported"] is False


def test_local_pptx_parser_preserves_slide_and_shape_locations():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    title.name = "Title"
    title.text = "设备维护"
    body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    body.name = "Body"
    body.text = "检查温度。\n记录结果。"
    table_shape = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1))
    table_shape.table.cell(0, 0).text = "项目"
    table_shape.table.cell(0, 1).text = "限值"
    table_shape.table.cell(1, 0).text = "温度"
    table_shape.table.cell(1, 1).text = "80℃"
    stream = io.BytesIO()
    presentation.save(stream)

    parsed = parse_local("设备维护.pptx", stream.getvalue(), "pptx-rev")
    assert parsed["parser"] == "local-baseline@2"
    assert [b["slide"] for b in parsed["blocks"]] == [1, 1, 1]
    assert [b["shape_name"] for b in parsed["blocks"]] == ["Title", "Body", table_shape.name]
    assert parsed["blocks"][1]["text"] == "检查温度。\n记录结果。"
    assert parsed["blocks"][2]["type"] == "table"
    assert "温度 | 80℃" in parsed["blocks"][2]["text"]
    assert parsed["slide_size"]["width"] > 0
    assert parsed["blocks"][1]["position"]["width"] == int(body.width)


def test_malformed_pptx_is_rejected_with_supported_format_message():
    with pytest.raises(ValueError, match="PPTX"):
        parse_local("broken.pptx", b"not-a-presentation", "rev")


def test_mineu_can_be_added_through_the_parser_adapter_contract(tmp_path):
    class MineUAdapter:
        version = "mineu@contract-fixture"

        def parse(self, name, data, revision):
            assert data.read() == b"source"
            return {"schema_version": 1, "revision": revision, "parser": self.version,
                    "blocks": [{"id": "m0", "type": "text", "text": "normalized source",
                                "page": None, "section_path": []}], "warnings": []}

    settings = Settings(database_url=f"sqlite:///{tmp_path / 'db'}", storage=tmp_path, vector_dim=3,
                        api_token="admin", max_file_bytes=100)
    lab = Lab(settings, embedder=object(), parser_adapters={"mineu": MineUAdapter()})
    kb = lab.create_kb("local", "mineu")
    doc = lab.ingest("local", kb["id"], "source.txt", io.BytesIO(b"source"), "manual", "mineu-doc")
    parsed = lab.parsed("local", doc["id"], "mineu")
    assert parsed["parser"] == "mineu@contract-fixture"
    assert parsed["blocks"][0]["id"] == "m0"
