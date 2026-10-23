import io

from docx import Document

from raglab.parsing import parse_local


def test_docx_preserves_table_order():
    doc = Document()
    doc.add_heading("工艺", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "部件"
    table.cell(0, 1).text = "温度"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "80"
    doc.add_paragraph("结束")
    stream = io.BytesIO()
    doc.save(stream)
    parsed = parse_local("test.docx", stream.getvalue(), "r")
    assert [block["type"] for block in parsed["blocks"]] == ["text", "table", "text"]
    assert "A | 80" in parsed["blocks"][1]["text"]
    assert parsed["parser"] == "local-baseline@2"
