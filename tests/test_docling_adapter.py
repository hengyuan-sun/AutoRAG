import hashlib
import os
import time
import tracemalloc
from importlib.metadata import version

import pytest
from docx import Document
from pptx import Presentation
from pydantic import ValidationError
from pypdf import PdfWriter

from raglab.document_ir.docling_adapter import DoclingOptions, DoclingParserAdapter
from raglab.document_ir.models import ItemKind


def source_docx(path):
    document = Document()
    document.add_heading("维修规范", level=1)
    document.add_paragraph("检查阀门并记录状态。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "限值"
    table.cell(1, 0).text = "压力"
    table.cell(1, 1).text = "0.8MPa"
    document.save(path)
    return path.read_bytes()


def test_docling_adapter_maps_heading_paragraph_table_and_provenance(tmp_path):
    path = tmp_path / "fixture.docx"
    source = source_docx(path)
    ir = DoclingParserAdapter().parse(path.name, path, hashlib.sha256(source).hexdigest())

    assert [item.type for item in ir.items] == [ItemKind.HEADING.value, ItemKind.PARAGRAPH.value,
                                                ItemKind.TABLE.value]
    assert ir.items[0].text == "维修规范"
    assert ir.items[1].text == "检查阀门并记录状态。"
    assert ir.items[2].cells[3].text == "0.8MPa"
    assert all(item.uid and item.provenance.source == "docling" for item in ir.items)
    assert ir.meta.sha256 == hashlib.sha256(source).hexdigest()


def test_docling_adapter_does_not_fabricate_missing_location(tmp_path):
    path = tmp_path / "fixture.docx"
    source_docx(path)
    ir = DoclingParserAdapter().parse(path.name, path, "fixture-revision")

    # DOCX text items have logical order but the simple Docling Word backend does not paginate them.
    assert all(location.page is None and location.bbox is None
               for item in ir.items for location in item.provenance.locations)
    assert all(item.provenance.locations == [] for item in ir.items)


def test_docling_identity_includes_package_and_pipeline_config():
    default = DoclingParserAdapter(DoclingOptions())
    alternate = DoclingParserAdapter(DoclingOptions(do_ocr=True))

    assert version("docling-slim") in default.version
    assert default.version != alternate.version
    assert "config=" in default.version


def test_docling_runtime_cannot_disable_offline_mode():
    with pytest.raises(ValidationError, match="offline"):
        DoclingOptions(offline=False)


def test_pdf_without_local_model_runtime_fails_with_install_instruction(tmp_path):
    artifacts = tmp_path / "models"
    artifacts.mkdir()
    (artifacts / "placeholder").write_text("offline asset marker")
    pdf_path = tmp_path / "placeholder.pdf"
    pdf_path.write_bytes(b"pdf fixture")

    with pytest.raises(ValueError, match=r"docling-pdf"):
        DoclingParserAdapter(DoclingOptions(artifacts_path=artifacts)).parse(
            pdf_path.name, pdf_path, "pdf-revision")


def test_docling_parser_translates_invalid_document_to_actionable_error(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a DOCX package")

    with pytest.raises(ValueError, match="Docling解析失败"):
        DoclingParserAdapter().parse(path.name, path, "broken-revision")


def test_docling_adapter_parses_pptx_and_plain_text_with_stable_uids(tmp_path):
    pptx_path = tmp_path / "fixture.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.add_textbox(0, 0, 1000, 500).text = "维护步骤"
    presentation.save(pptx_path)
    text_path = tmp_path / "plain.txt"
    text_path.write_text("第一行\n第二行", encoding="utf-8")
    adapter = DoclingParserAdapter()

    pptx_ir = adapter.parse(pptx_path.name, pptx_path, "pptx-revision")
    text_ir = adapter.parse(text_path.name, text_path, "text-revision")
    repeated = adapter.parse(pptx_path.name, pptx_path, "pptx-revision")

    assert any(item.text == "维护步骤" for item in pptx_ir.items)
    assert any("第一行" in item.text for item in text_ir.items)
    assert [item.uid for item in repeated.items] == [item.uid for item in pptx_ir.items]


@pytest.mark.integration
def test_docling_pdf_page_order_and_resource_report(tmp_path, capsys):
    artifact_path = os.environ.get("RAGLAB_DOCLING_ARTIFACTS_PATH")
    if not artifact_path:
        pytest.skip("set RAGLAB_DOCLING_ARTIFACTS_PATH after offline Docling PDF models are prepared")
    pdf_path = tmp_path / "thousand-pages.pdf"
    writer = PdfWriter()
    for _ in range(1000):
        writer.add_blank_page(width=612, height=792)
    with pdf_path.open("wb") as stream:
        writer.write(stream)

    adapter = DoclingParserAdapter(DoclingOptions(artifacts_path=artifact_path, max_pages=1000))
    tracemalloc.start()
    started = time.perf_counter()
    ir = adapter.parse(pdf_path.name, pdf_path, "large-pdf-fixture")
    elapsed = time.perf_counter() - started
    _current, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    pages = [page.page for page in ir.pages]
    assert pages == list(range(1, 1001))
    print(f"Docling 1000-page PDF: elapsed_s={elapsed:.2f}, python_peak_bytes={peak}")
